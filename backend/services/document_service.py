import os
import io
import re
from typing import List, Tuple
from pypdf import PdfReader
import docx

from backend.core.config import get_settings
from backend.core.security import sanitize_filename

settings = get_settings()


class DocumentService:
    """
    Handles text extraction from PDF, DOCX, and TXT files,
    sliding-window text chunking, and document generation.
    """

    @staticmethod
    def extract_text_from_bytes(file_bytes: bytes, filename: str) -> str:
        """
        Extracts raw textual content from uploaded file bytes based on extension.
        """
        _, ext = os.path.splitext(filename.lower())

        if ext == ".pdf":
            try:
                reader = PdfReader(io.BytesIO(file_bytes))
                pages_text = []
                for i, page in enumerate(reader.pages):
                    text = page.extract_text() or ""
                    if text.strip():
                        pages_text.append(f"--- Page {i + 1} ---\n{text.strip()}")
                return "\n\n".join(pages_text).strip()
            except Exception as e:
                raise ValueError(f"Failed to extract text from PDF: {str(e)}")

        elif ext == ".docx":
            try:
                doc = docx.Document(io.BytesIO(file_bytes))
                paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
                return "\n\n".join(paragraphs).strip()
            except Exception as e:
                raise ValueError(f"Failed to extract text from DOCX: {str(e)}")

        elif ext == ".txt":
            try:
                return file_bytes.decode("utf-8")
            except UnicodeDecodeError:
                # Fallback to latin-1
                return file_bytes.decode("latin-1", errors="replace")

        else:
            raise ValueError(f"Unsupported document format: '{ext}'. Supported formats: .pdf, .docx, .txt")

    @staticmethod
    def chunk_text(text: str, chunk_size: int = 500, overlap: int = 100) -> List[str]:
        """
        Splits text into overlapping chunks of approximately chunk_size characters,
        attempting to break at sentence or word boundaries.
        """
        text = text.strip()
        if not text:
            return []

        if len(text) <= chunk_size:
            return [text]

        chunks: List[str] = []
        start = 0
        total_len = len(text)

        while start < total_len:
            end = start + chunk_size

            if end >= total_len:
                chunk = text[start:].strip()
                if chunk:
                    chunks.append(chunk)
                break

            # Try to break at paragraph, newline, or punctuation
            breakpoint = -1
            for delimiter in ["\n\n", "\n", ". ", "? ", "! ", " "]:
                pos = text.rfind(delimiter, start, end)
                if pos != -1 and pos > start + (chunk_size // 2):
                    breakpoint = pos + len(delimiter)
                    break

            if breakpoint == -1:
                breakpoint = end

            chunk = text[start:breakpoint].strip()
            if chunk:
                chunks.append(chunk)

            # Advance start by chunk length minus overlap
            start = max(start + 1, breakpoint - overlap)

        return chunks

    @staticmethod
    def save_uploaded_file(file_bytes: bytes, filename: str) -> str:
        """
        Saves uploaded file bytes to local storage directory securely.
        Returns the relative file path.
        """
        os.makedirs(settings.STORAGE_DIR, exist_ok=True)
        sanitized = sanitize_filename(filename)
        dest_path = os.path.join(settings.STORAGE_DIR, sanitized)

        # Handle duplicate filenames by appending counter
        base, ext = os.path.splitext(dest_path)
        counter = 1
        while os.path.exists(dest_path):
            dest_path = f"{base}_{counter}{ext}"
            counter += 1

        with open(dest_path, "wb") as f:
            f.write(file_bytes)

        return dest_path

    @staticmethod
    def create_document_file(title: str, content: str, file_format: str = "txt") -> Tuple[str, str]:
        """
        Creates a new downloadable file (.txt or .docx) from generated content.
        Returns (relative_file_path, filename).
        """
        gen_dir = os.path.join(settings.STORAGE_DIR, "generated")
        os.makedirs(gen_dir, exist_ok=True)

        clean_title = sanitize_filename(title)
        base_name, _ = os.path.splitext(clean_title)

        if file_format.lower() == "docx":
            out_filename = f"{base_name}.docx"
            out_path = os.path.join(gen_dir, out_filename)
            doc = docx.Document()
            doc.add_heading(base_name.replace("_", " ").title(), level=1)
            for paragraph in content.split("\n\n"):
                if paragraph.strip():
                    doc.add_paragraph(paragraph.strip())
            doc.save(out_path)
            return out_path, out_filename
        else:
            out_filename = f"{base_name}.txt"
            out_path = os.path.join(gen_dir, out_filename)
            with open(out_path, "w", encoding="utf-8") as f:
                f.write(content)
            return out_path, out_filename
